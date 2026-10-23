CELL1
# ============================================================
# 1. SETUP
# ============================================================
import sys, subprocess

def _pip_install(pkgs):
    subprocess.run([sys.executable, "-m", "pip", "install", "-q", *pkgs], check=True)

for pkg, import_name in [("openpyxl", "openpyxl"), ("xlsxwriter", "xlsxwriter"), ("statsmodels", "statsmodels")]:
    try:
        __import__(import_name)
    except ImportError:
        _pip_install([pkg])

import re, json, base64, hashlib, warnings
from collections import OrderedDict, defaultdict
import numpy as np
import pandas as pd
import requests
import openpyxl
from statsmodels.stats.proportion import proportion_confint

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

IN_COLAB = "google.colab" in sys.modules
print("Running in Google Colab:", IN_COLAB)
print("pandas", pd.__version__, "| numpy", np.__version__)


CELL2
# ============================================================
# 2. CONFIG
# ============================================================
CONFIG = {
    "gbm_url": ("https://raw.githubusercontent.com/omixer/omixer-rpmR/"
                "184f1cc99aefed722e20eb00eda082348a064c4e/inst/extdata/GBMs.v1.0.txt"),
    "output_xlsx": "GBM_completeness_and_statistics_v2.xlsx",
    "pangenome_thresholds": {"core": 99.0, "soft_core": 95.0, "shell": 15.0},

    # Column-name candidates (case-insensitive), so small header variations don't break the notebook.
    "ko_column_candidates": ["KO IDs", "KO_ID", "KO_IDs", "KO", "KEGG_ko"],
    "metadata_header_candidates": ["gene", "non-unique gene name", "annotation", "ko ids", "ko_id",
                                    "ko", "count_present", "freq"],
}
print("Config loaded.")


CELL3
# ============================================================
# 3. OFFICIAL GBM DATABASE: download with verified offline fallback
# ============================================================
_GBM_EMBEDDED_B64 = (
    "TUdCMDAxCVNlcm90b25pbiBzeW50aGVzaXMgSQpDT0cyNTA5CkswMTU5MwlDT0cwMDc2Ci8vLwpNR0IwMDIJU2Vyb3RvbmluIHN5bnRoZXNpcyBJSQpLMDE1OTMJQ09HMDA3NgpLMDA0OTIKLy8vCk1HQjAwMwlNZWxhdG9uaW4gc3ludGhlc2lzCmJhY3ROT0cyNjMxNglwcm9OT0czMDE5MQpLMDA1NDMJQ09HMDUwMAovLy8KTUdCMDA0CUt5bnVyZW5pbmUgc3ludGhlc2lzCkswMDQ1MwlLMDA0NjMJYmFjdE5PRzEyNDUzCUNPRzM0ODMKSzAxNDMyCUswNzEzMAlDT0cxODc4Ci8vLwpNR0IwMDUJVHJ5cHRvcGhhbiBzeW50aGVzaXMKSzAxNjU3LEswMTY1OAlLMTM0OTcJSzEzNTAxCUsxMzUwMwpLMDA3NjYJSzEzNDk3CkswMTgxNwlLMTM0OTgJSzEzNTAxCkswMTYwOQlLMTM0OTgJSzEzNTAxCkswMTY5NSxLMDE2OTYJSzAxNjk1LEswNjAwMQovLy8KTUdCMDA2CUdsdXRhbWF0ZSBzeW50aGVzaXMgSQpLMDAyNjEJSzAwMjYyCi8vLwpNR0IwMDcJR2x1dGFtYXRlIHN5bnRoZXNpcyBJSQpLMDAyNjUsSzAwMjY2CUswMDI4NAovLy8KTUdCMDA4CUFjZXR5bGNob2xpbmUgc3ludGhlc2lzCkswMDYyNAliYWN0Tk9HODI2MTcJYmFjdE5PRzAyMTU5Ci8vLwpNR0IwMDkJSGlzdGFtaW5lIHN5bnRoZXNpcwpLMDE1OTAJVElHUjAwNTQxCWJhY3ROT0cwNTgyOQliYWN0Tk9HMTY2MjcKLy8vCk1HQjAxMAlIaXN0YW1pbmUgZGVncmFkYXRpb24KSzAwMjc2CWJhY3ROT0cwODM0MAliYWN0Tk9HMDE2NjIKLy8vCk1HQjAxMQlDb3J0aXNvbCBkZWdyYWRhdGlvbgpDT0czOTU5CmJhY3ROT0cwMTg0NApDT0cxMDYzCmJhY3ROT0cwMzI2OQovLy8KTUdCMDEyCURvcGFtaW5lIHN5bnRoZXNpcwpLMDA1MDUJVElHUjAxMjY5CWJhY3ROT0cxNDA5OAliYWN0Tk9HMjAyNDkKSzAxNTkzCUswMTU5MglUSUdSMDM4MTEJVElHUjAzODEyCUNPRzAwNzYKLy8vCk1HQjAxMwlOb3JhZHJlbmFsaW5lIHN5bnRoZXNpcwpOT0cxMzI1NTMKLy8vCk1HQjAxNAlBZHJlbmFsaW5lIHN5bnRoZXNpcwpOT0cxMzM2NjMKLy8vCk1HQjAxNQlwLUNyZXNvbCBzeW50aGVzaXMKSzE4NDI3LEsxODQyOAlLMTU3NjAsSzE1NzYxLEsxNTc2MixLMTU3NjMsSzE1NzY0CUswMzE1MAliYWN0Tk9HMjg4NzYKLy8vCk1HQjAxNglwLUNyZXNvbCBkZWdyYWRhdGlvbgpLMDU3OTcJSzAzMzgwCUswMDQ0OCxLMDA0NDkJVElHUjAyNDIzLFRJR1IwMjQyMglLMDA0NDgsVElHUjAyNDIyCVRJR1IwMjQyMyxLMDA0NDkKLy8vCk1HQjAxNwlJUEEgc3ludGhlc2lzIChDbG9zdHJpZGl1bSBzcG9yb2dlbmVzIHBhdGh3YXkpCmZpcm1OT0cwNDI5MApmaXJtTk9HMDA2MjYsYmFjdE5PRzA3ODgxLGJhY3ROT0cxNTM0MQovLy8KTUdCMDE4CUt5bnVyZW5pbmUgZGVncmFkYXRpb24KSzAwNDg2CWJhY3ROT0cwMjU5MgovLy8KTUdCMDE5CUdBQkEgZGVncmFkYXRpb24KSzA3MjUwCUsxNDI2OAlLMDA4MjMJSzEzNTI0CUswMDYxMwlLMTY4NzEKLy8vCk1HQjAyMAlHQUJBIHN5bnRoZXNpcyBJCkswOTQ3MApLMDk0NzEKSzA5NDcyCkswOTQ3MwovLy8KTUdCMDIxCUdBQkEgc3ludGhlc2lzIElJCkswOTI1MQpLMDAxMzcKLy8vCk1HQjAyMglHQUJBIHN5bnRoZXNpcyBJSUkKSzAxNTgwCWJhY3ROT0cwMTc5MwovLy8KTUdCMDIzCURvcGFtaW5lIGRlZ3JhZGF0aW9uCkswMTAxNAlLMDA1NDUJSzAwMjc0CUswMDI3NglLMTMzNzEsSzEzMzcyCWJhY3ROT0cwODM0MAovLy8KTUdCMDI0CURPUEFDIHN5bnRoZXNpcwpLMDAyNzQJSzAwMjc2CUsxMzM3MSxLMTMzNzIJYmFjdE5PRzA4MzQwCkswMDEyOAlLMDAxMjkJQ09HMTAxMgovLy8KTUdCMDI1CU5pdHJpYyBveGlkZSBzeW50aGVzaXMgSSAoTk8gc3ludGhhc2UpCkswMDQ5MQliYWN0Tk9HMTI4MzYJYmFjdE5PRzUxNjc4Ci8vLwpNR0IwMjYJTml0cmljIG94aWRlIHN5bnRoZXNpcyBJSSAobml0cml0ZSByZWR1Y3Rhc2UpCkswMDM2OAlLMTU4NjQJYmFjdE5PRzEwMjQ3CUNPRzMzMDMKLy8vCk1HQjAyNwlOaXRyaWMgb3hpZGUgZGVncmFkYXRpb24gSSAoTk8gZGlveHlnZW5hc2UpCkswNTkxNgliYWN0Tk9HMDUwNTYJYmFjdE5PRzM2OTg0Ci8vLwpNR0IwMjgJTml0cmljIG94aWRlIGRlZ3JhZGF0aW9uIElJIChOTyByZWR1Y3Rhc2UpCkswNDU2MSxLMDIzMDUJYmFjdE5PRzAzNzY2LGJhY3ROT0cxODE5MglLMDQ1NjEsYmFjdE5PRzE4MTkyCWJhY3ROT0cwMzc2NixLMDIzMDUKLy8vCk1HQjAyOQlDbHBCIChBVFAtZGVwZW5kZW50IGNoYXBlcm9uZSBwcm90ZWluKQpLMDM2OTUJVElHUjAzMzQ2CUNPRzA1NDIKLy8vCk1HQjAzMAlQb2x5c2FjY2hhcmlkZSBBCmJhY3ROT0cwMDAxNApiYWN0Tk9HMTg1MTkKYmFjdE5PRzI0NzExCWJhY3ROT0cxODYzMAliYWN0Tk9HMzExMTYKLy8vCk1HQjAzMQkxNy1iZXRhLUVzdHJhZGlvbCBkZWdyYWRhdGlvbgpDT0cwMzAwCUswNzEyNAovLy8KTUdCMDMyCVF1aW5vbGluaWMgYWNpZCBzeW50aGVzaXMKSzAzNTE3LEswMDI3OAlUSUdSMDA1NTAsVElHUjAwNTUxCWJhY3ROT0cwMjIzOSxDT0cwMDI5CUswMzUxNyxUSUdSMDA1NTEJSzAzNTE3LENPRzAwMjkJVElHUjAwNTUwLEswMDI3OAlUSUdSMDA1NTAsQ09HMDAyOQliYWN0Tk9HMDIyMzksSzAwMjc4CWJhY3ROT0cwMjIzOSxUSUdSMDA1NTEKLy8vCk1HQjAzMwlRdWlub2xpbmljIGFjaWQgZGVncmFkYXRpb24KSzAwNzY3CVRJR1IwMDA3OAliYWN0Tk9HMDA0NDEKSzAwOTY5CWJhY3ROT0czMDQ2NAliYWN0Tk9HNjkzMzUKSzAxOTE2CUswMTk1MAlUSUdSMDA1NTIJQ09HMDE3MQovLy8KTUdCMDM0CUlzb3ZhbGVyaWMgYWNpZCBzeW50aGVzaXMgSSAoS0FESCBwYXRod2F5KQpLMDA4MjYJVElHUjAxMTIyCVRJR1IwMTEyMwlDT0cwMTE1CkswMDE2NixLMDAxNjcJSzExMzgxCkswOTY5OQpLMDAzODIJVElHUjAxMzUwCkswMTA3MwlLMTA4MDYJYmFjdE5PRzI3MDQzCi8vLwpNR0IwMzUJSXNvdmFsZXJpYyBhY2lkIHN5bnRoZXNpcyBJSSAoS0FEQyBwYXRod2F5KQpLMDA4MjYJVElHUjAxMTIyCVRJR1IwMTEyMwlDT0cwMTE1CkswMTU2OAlDT0czOTYxCkswMDEyOAliYWN0Tk9HMDA4NTIKLy8vCk1HQjAzNglTLUFkZW5vc3lsbWV0aGlvbmluZSAoU0FNKSBzeW50aGVzaXMKSzAwNzg5CVRJR1IwMTAzNAliYWN0Tk9HMDI5MzcKLy8vCk1HQjAzNwlJbm9zaXRvbCBzeW50aGVzaXMKSzAxODU4CWJhY3ROT0cwNjgxMQpLMDEwOTIJSzE4NjQ5CUNPRzA0ODMKLy8vCk1HQjAzOAlJbm9zaXRvbCBkZWdyYWRhdGlvbgpLMDAwMTAJVElHUjA0MzgwCWJhY3ROT0cwNzE1MwovLy8KTUdCMDM5CWctSHlkcm94eWJ1dHlyaWMgYWNpZCAoR0hCKSBkZWdyYWRhdGlvbgpLMDAwNDMJSzA4MzE4CUsxODEyMAliYWN0Tk9HMTA1NjUKLy8vCk1HQjA0MAlNZW5hcXVpbm9uZSBzeW50aGVzaXMgKHZpdGFtaW4gSzIpIEkKSzAxODUxCUswMjM2MQlLMDI1NTIJSzA0NzgxCkswMjU1MQpLMDg2ODAKSzAyNTQ5CkswMTkxMQpLMDE2NjEKSzEyMDczCUsxOTIyMgpLMDI1MjMJSzAwODA1CkswMjU0OApLMDMxODMKLy8vCk1HQjA0MQlNZW5hcXVpbm9uZSBzeW50aGVzaXMgKHZpdGFtaW4gSzIpIElJIChhbHRlcm5hdGl2ZSBwYXRod2F5OiBmdXRhbG9zaW5lIHBhdGh3YXkpCksxMTc4MgpLMTE3ODMKSzExNzg0CksxMTc4NQpLMDI1NDgKSzAzMTgzCi8vLwpNR0IwNDIJUFVGQXMgc3ludGhlc2lzIChBQSwgRVBBLCBESEEpCmJhY3ROT0cwNDU3OCxiYWN0Tk9HNjIyNDUsYmFjdE5PRzAyMTk0LGJhY3ROT0c0ODMwNwlDT0czMzIxLGJhY3ROT0cwMjE5NCxiYWN0Tk9HNDgzMDcKLy8vCk1HQjA0MwlBY2V0YXRlIHN5bnRoZXNpcyBJCkswMDY1NglLMDAxNjksSzAwMTcwLEswMDE3MSxLMDAxNzIJSzAzNzM3CkswMDYyNQlLMTM3ODgJSzE1MDI0CkswMDkyNQovLy8KTUdCMDQ0CUFjZXRhdGUgc3ludGhlc2lzIElJCkswMDY1NglLMDAxNjksSzAwMTcwLEswMDE3MSxLMDAxNzIJSzAzNzM3CkswMTkwNQovLy8KTUdCMDQ1CUFjZXRhdGUgc3ludGhlc2lzIElJSQpLMDA2NTYJSzAwMTY5LEswMDE3MCxLMDAxNzEsSzAwMTcyCUswMzczNwpLMTgxMTgKLy8vCk1HQjA0NglBY2V0YXRlIHN5bnRoZXNpcyBJVgpLMDE1NjgKSzE0NTE5Ci8vLwpNR0IwNDcJQWNldGF0ZSBkZWdyYWRhdGlvbgpLMDE4OTUJYmFjdE5PRzAwOTQ3CWJhY3ROT0cwMTI1OQovLy8KTUdCMDQ4CVByb3Bpb25hdGUgc3ludGhlc2lzIEkKSzEzOTIzCkswMDkzMglLMTk2OTcKLy8vCk1HQjA0OQlUcnlwdG9waGFuIGRlZ3JhZGF0aW9uCkswMTY2NwlUSUdSMDI2MTcJYmFjdE5PRzAyODgxCi8vLwpNR0IwNTAJR2x1dGFtYXRlIGRlZ3JhZGF0aW9uIEkKSzAwMjYwCUsxNTM3MQlLMDAyNjEKSzAwMTA5CkswMTAzOSxLMDEwNDAKSzAxNjE1Ci8vLwpNR0IwNTEJR2x1dGFtYXRlIGRlZ3JhZGF0aW9uIElJCkswMTg0NixLMTkyNjgKSzA0ODM1Ci8vLwpNR0IwNTIJQnV0eXJhdGUgc3ludGhlc2lzIEkKSzAwMjQ4CkswMDYzNApLMDA5MjkKLy8vCk1HQjA1MwlCdXR5cmF0ZSBzeW50aGVzaXMgSUkKSzAwMjQ4CkswMTAzNCxLMDEwMzUJSzE5NzA5Ci8vLwpNR0IwNTQJUHJvcGlvbmF0ZSBzeW50aGVzaXMgSUkKSzAxMDI2Ci8vLwpNR0IwNTUJUHJvcGlvbmF0ZSBzeW50aGVzaXMgSUlJCkswMTg0NwlLMDE4NDksSzAxODQ4CUsxMTk0MgpLMDU2MDYKSzAzNDE2LEsxNzQ4OSxLMTc0OTAKLy8vCk1HQjA1NglQcm9waW9uYXRlIGRlZ3JhZGF0aW9uIEkKSzAxOTA4CkswMTY1OQpLMDE3MjAKSzAxNjgyCkswMzQxNwo="
)

def _load_gbm_text():
    embedded_bytes = base64.b64decode(_GBM_EMBEDDED_B64)
    embedded_sha = hashlib.sha256(embedded_bytes).hexdigest()
    downloaded_bytes = None
    try:
        resp = requests.get(CONFIG["gbm_url"], timeout=15)
        resp.raise_for_status()
        downloaded_bytes = resp.content
    except Exception as e:
        print(f"[info] Could not download GBMs.v1.0.txt ({e}); using embedded offline copy.")
    if downloaded_bytes is not None:
        downloaded_sha = hashlib.sha256(downloaded_bytes).hexdigest()
        if downloaded_sha == embedded_sha:
            print("[ok] Downloaded GBMs.v1.0.txt is byte-identical to the pinned embedded copy.")
            print(f"     SHA-256: {downloaded_sha}")
            return downloaded_bytes.decode("utf-8"), downloaded_sha, "downloaded (verified)"
        print("[warning] Downloaded file does not match the pinned SHA-256 — using the verified "
              "embedded copy instead, to guarantee this run matches prior runs.")
    print(f"[ok] Using embedded GBMs.v1.0.txt. SHA-256: {embedded_sha}")
    return embedded_bytes.decode("utf-8"), embedded_sha, "embedded fallback"

GBM_RAW_TEXT, GBM_SHA256, GBM_SOURCE_MODE = _load_gbm_text()
print("GBM source used for this run:", GBM_SOURCE_MODE)


CELL4
# ============================================================
# 4. PARSE OFFICIAL GBM STEP LOGIC
# ============================================================
def parse_gbm_db(raw_text):
    # Each module block ends in a line containing only '///'.
    # First line of a block: 'MGBxxx<TAB>Module name'.
    # Every subsequent line = one obligatory functional step (steps are AND-ed).
    # Within a step-line: TAB-separated fields = OR alternatives.
    # Within one field: COMMA-separated ortholog IDs = AND-required complex subunits.
    modules = OrderedDict()
    blocks = [b for b in raw_text.strip().split("///") if b.strip()]
    for b in blocks:
        lines = [l for l in b.strip("\n").split("\n") if l.strip() != ""]
        if not lines:
            continue
        header = lines[0].split("\t")
        mod_id = header[0].strip()
        mod_name = header[1].strip() if len(header) > 1 else ""
        steps = []
        for line in lines[1:]:
            alt_groups = [[m.strip() for m in field.split(",") if m.strip()] for field in line.split("\t")]
            alt_groups = [g for g in alt_groups if g]
            if alt_groups:
                steps.append(alt_groups)
        if mod_id:
            modules[mod_id] = {"name": mod_name, "steps": steps}
    return modules

GBM = parse_gbm_db(GBM_RAW_TEXT)
assert len(GBM) == 56, f"Expected 56 GBM modules, parsed {len(GBM)} — GBM source file may be corrupted."
assert GBM["MGB005"]["name"] == "Tryptophan synthesis" and len(GBM["MGB005"]["steps"]) == 5
assert GBM["MGB022"]["steps"][0][0] == ["K01580"]
print(f"[ok] Parsed {len(GBM)} GBM modules; sanity checks passed.")

ALL_GBM_KOS = set()
ko_to_module_steps = defaultdict(list)
for mid, d in GBM.items():
    for s_i, step in enumerate(d["steps"], 1):
        for grp in step:
            for m in grp:
                if re.fullmatch(r"K\d{5}", m):
                    ALL_GBM_KOS.add(m)
                    ko_to_module_steps[m].append((mid, s_i))
print(f"[ok] {len(ALL_GBM_KOS)} distinct KOs referenced across the 56 official modules.")


CELL5
# ============================================================
# 5. UPLOAD YOUR TABLE 8 FILE
# ============================================================
if IN_COLAB:
    from google.colab import files
    print("Select your Table 8 .xlsx file...")
    uploaded = files.upload()
    assert len(uploaded) >= 1, "No file uploaded."
    table8_path = list(uploaded.keys())[0]
else:
    table8_path = "Table_8.xlsx"

print("Using Table 8 file:", table8_path)


CELL6
# ============================================================
# 6. PARSE TABLE 8 (KO-annotated format)
# ============================================================
def parse_table8_with_ko(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    max_row, max_col = ws.max_row, ws.max_column
    header = [ws.cell(row=2, column=c).value for c in range(1, max_col + 1)]

    metadata_set = set(CONFIG["metadata_header_candidates"])
    ko_col = None
    for i, h in enumerate(header, start=1):
        if h is not None and str(h).strip().lower() in [c.lower() for c in CONFIG["ko_column_candidates"]]:
            ko_col = i
            break
    assert ko_col is not None, (
        f"Could not find a KO-ID column. Expected one of {CONFIG['ko_column_candidates']} in row 2. "
        f"Columns found: {header}. Add your column's exact name to CONFIG['ko_column_candidates'] "
        f"and re-run from this cell."
    )
    print(f"[ok] KO column detected at column {ko_col} ('{header[ko_col-1]}').")

    genome_start_col = None
    for i, h in enumerate(header, start=1):
        if h is not None and str(h).strip().lower() not in metadata_set and i > ko_col:
            genome_start_col = i
            break
    assert genome_start_col is not None, (
        "Could not locate where genome ID columns start (looked for the first non-metadata column "
        "after the KO column). Check that row 2 is the header row and genome IDs follow the KO column."
    )
    genome_cols, genome_ids = [], []
    for c in range(genome_start_col, max_col + 1):
        h = ws.cell(row=2, column=c).value
        if h is not None:
            genome_cols.append(c)
            genome_ids.append(h)
    n_genomes = len(genome_ids)
    assert n_genomes > 0, "No genome columns detected."
    print(f"[ok] Detected {n_genomes} genome columns (columns {genome_cols[0]}-{genome_cols[-1]}).")

    count_col = None
    for i, h in enumerate(header, start=1):
        if h is not None and str(h).strip().lower() == "count_present":
            count_col = i
            break

    records = []
    current_section = None
    for r in range(3, max_row + 1):
        a = ws.cell(row=r, column=1).value
        ko_val = ws.cell(row=r, column=ko_col).value
        count_val = ws.cell(row=r, column=count_col).value if count_col else None
        if a is None:
            continue
        # Section header rows have no KO id / no count_present.
        if ko_val is None and count_val is None:
            current_section = str(a).strip()
            continue
        presence = [ws.cell(row=r, column=col).value for col in genome_cols]
        presence = [int(v) if v is not None else 0 for v in presence]
        # A cell may carry more than one KO, comma-separated (e.g. a Panaroo cluster whose gene
        # name bundles an unfused gene and its fusion-protein counterpart). Every listed KO is
        # treated as potentially represented by this row: the row's presence/absence vector is
        # shared across all of them, since Table 8 gives no finer per-KO resolution than per-row.
        ko_list = [k.strip() for k in str(ko_val).split(",") if k.strip()] if ko_val else []
        records.append({
            "row": r, "section": current_section, "Gene": a,
            "KO_list": ko_list,
            "presence": presence,
        })

    df = pd.DataFrame(records)
    assert len(df) > 0, "No gene rows were parsed — check the file structure."
    n_no_ko = sum(1 for l in df["KO_list"] if len(l) == 0)
    n_multi_ko = sum(1 for l in df["KO_list"] if len(l) > 1)
    if n_no_ko:
        print(f"[info] {n_no_ko} gene row(s) have no KO ID and will be excluded from GBM scoring "
              f"(shown in the Unassigned_Rows sheet of the output).")
    if n_multi_ko:
        print(f"[info] {n_multi_ko} gene row(s) carry more than one comma-separated KO ID; each is "
              f"registered as potential evidence for that row's presence/absence pattern.")
    return df, genome_ids

df8, genome_ids = parse_table8_with_ko(table8_path)
N_GENOMES = len(genome_ids)
print(f"[ok] Parsed {len(df8)} gene rows across {df8['section'].nunique()} sections, "
      f"{N_GENOMES} genomes.")



CELL7
# ============================================================
# 7. DIRECT KO -> OFFICIAL GBM MODULE/STEP LOOKUP  (no symbol guessing)
# ============================================================
presence_matrix = np.array(df8["presence"].tolist())
row_to_idx = {row: i for i, row in enumerate(df8["row"])}

# Flat (row, KO) pairs — the natural unit for both the lookup index and the catalog-membership report,
# since a single row can now carry more than one KO.
pair_records = []
for _, r in df8.iterrows():
    for ko in r["KO_list"]:
        pair_records.append({"row": r["row"], "section": r["section"], "Gene": r["Gene"], "KO": ko})
ko_row_pairs = pd.DataFrame(pair_records)

ko_to_rows = defaultdict(set)
for _, p in ko_row_pairs.iterrows():
    ko_to_rows[p["KO"]].add(p["row"])

def ko_vec(ko):
    rows = ko_to_rows.get(ko, set())
    if not rows:
        return None
    idxs = [row_to_idx[r] for r in rows]
    return (presence_matrix[idxs, :].sum(axis=0) > 0).astype(int)

# Flag individual (row, KO) pairs whose KO is not part of ANY official GBM module — reported per KO,
# not per row, so a row with one valid and one invalid KO is not incorrectly cleared or incorrectly
# blanket-flagged.
ko_row_pairs["in_official_GBM_catalog"] = ko_row_pairs["KO"].isin(ALL_GBM_KOS)
unmatched = ko_row_pairs[~ko_row_pairs["in_official_GBM_catalog"]]
if len(unmatched) > 0:
    print(f"[WARNING] {len(unmatched)} (row, KO) pair(s) use a KO that is NOT part of any official GBM "
          f"module. This does not stop the analysis (that specific KO is simply not usable for GBM "
          f"scoring), but you should double-check these assignments are what you intended:")
    print(unmatched[["row", "section", "Gene", "KO"]].to_string(index=False))
else:
    print("[ok] Every assigned KO in this table is part of the official GBM catalog.")



CELL8
# ============================================================
# 8. DETERMINE WHICH MODULES ARE FULLY RECONSTRUCTABLE (fully automatic)
# ============================================================
step_coverage = {}
for mid, d in GBM.items():
    n_steps = len(d["steps"])
    n_with_data = 0
    missing_detail = []
    for s_i, step in enumerate(d["steps"], 1):
        has_full_alt = False
        for grp in step:
            ko_members = [m for m in grp if re.fullmatch(r"K\d{5}", m)]
            if ko_members and all(len(ko_to_rows.get(m, set())) > 0 for m in ko_members):
                has_full_alt = True
        n_with_data += int(has_full_alt)
        if not has_full_alt:
            missing_detail.append(s_i)
    step_coverage[mid] = {"n_steps": n_steps, "n_with_data": n_with_data,
                           "missing_steps": missing_detail,
                           "reconstructable": (n_with_data == n_steps and n_steps > 0)}

RECONSTRUCTABLE = sorted([m for m, v in step_coverage.items() if v["reconstructable"]])
NOT_RECONSTRUCTABLE = sorted([m for m, v in step_coverage.items()
                               if not v["reconstructable"] and v["n_with_data"] > 0])

print(f"[ok] {len(RECONSTRUCTABLE)} module(s) are fully reconstructable and will be quantitatively scored:")
for mid in RECONSTRUCTABLE:
    print(f"   {mid}  {GBM[mid]['name']}  ({step_coverage[mid]['n_steps']} step(s))")

print(f"\n[info] {len(NOT_RECONSTRUCTABLE)} module(s) have SOME but not all required step-level evidence "
      f"(excluded from quantitative scoring — see Excluded_or_Unresolved sheet):")
for mid in NOT_RECONSTRUCTABLE:
    v = step_coverage[mid]
    print(f"   {mid}  {GBM[mid]['name']}  ({v['n_with_data']}/{v['n_steps']} steps; "
          f"missing step(s) {v['missing_steps']})")



CELL9
# ============================================================
# 9. RESOLVE AND/OR COMPLETENESS PER GENOME
# ============================================================
comp_rows = []
step_rows_out = []
for mid in RECONSTRUCTABLE:
    steps = GBM[mid]["steps"]
    step_arrays = []
    for step in steps:
        group_arrays = []
        for grp in step:
            ko_members = [m for m in grp if re.fullmatch(r"K\d{5}", m)]
            vecs = [ko_vec(m) for m in ko_members]
            if not ko_members or any(v is None for v in vecs):
                group_arrays.append(np.zeros(N_GENOMES, dtype=int))
                continue
            garr = np.ones(N_GENOMES, dtype=int)
            for v in vecs:
                garr &= v
            group_arrays.append(garr)
        sarr = np.zeros(N_GENOMES, dtype=int)
        for g in group_arrays:
            sarr |= g
        step_arrays.append(sarr)
    smat = np.array(step_arrays)
    completeness = smat.sum(axis=0) / smat.shape[0]
    for s_i in range(smat.shape[0]):
        for gi in range(N_GENOMES):
            step_rows_out.append({"Module_ID": mid, "Module_name": GBM[mid]["name"],
                                   "Functional_step": s_i + 1, "Genome": genome_ids[gi],
                                   "Step_present": int(smat[s_i, gi])})
    for gi in range(N_GENOMES):
        comp_rows.append({"Module_ID": mid, "Module_name": GBM[mid]["name"], "Genome": genome_ids[gi],
                           "Completeness_%": completeness[gi] * 100})

comp_df = pd.DataFrame(comp_rows)
steps_df = pd.DataFrame(step_rows_out)
print(f"[ok] Computed per-genome completeness for {comp_df['Module_ID'].nunique()} modules "
      f"x {N_GENOMES} genomes.")


CELL10
# ============================================================
# 10. STATISTICS: exact 95% CI + pangenome core/accessory classification
# ============================================================
thr = CONFIG["pangenome_thresholds"]
def classify(pct):
    if pct >= thr["core"]: return "Core"
    if pct >= thr["soft_core"]: return "Soft-core"
    if pct >= thr["shell"]: return "Shell"
    return "Cloud"

ci_rows = []
for mid in RECONSTRUCTABLE:
    sub = comp_df[comp_df.Module_ID == mid]
    n_complete = int((sub["Completeness_%"] == 100).sum())
    ci_low, ci_high = proportion_confint(n_complete, N_GENOMES, alpha=0.05, method="beta")
    pct = 100 * n_complete / N_GENOMES
    ci_rows.append({
        "Module_ID": mid, "Module_name": GBM[mid]["name"], "N_required_steps": len(GBM[mid]["steps"]),
        "N_genomes_complete": n_complete, "N_genomes_total": N_GENOMES,
        "Prevalence_%": round(pct, 2), "CI95_low_%": round(ci_low * 100, 2), "CI95_high_%": round(ci_high * 100, 2),
        "Pangenome_category": classify(pct),
    })
prevalence_ci_df = pd.DataFrame(ci_rows).sort_values("Prevalence_%", ascending=False)
print("=== Module prevalence, exact 95% CI, and pangenome category ===")
print(prevalence_ci_df.to_string(index=False))

print("\n[note] No inferential test (t-test/ANOVA/chi-square/Fisher) is applied across genomes: the "
      "318 genomes are not independent sampling units without a phylogenetic correction, which was not "
      "supplied. Prevalence is reported descriptively with an exact confidence interval instead.")


CELL11
# ============================================================
# 11. WRITE THE FINAL EXCEL WORKBOOK
# ============================================================
excl_rows = [{"Module_ID": mid, "Module_name": GBM[mid]["name"],
              "N_required_steps": step_coverage[mid]["n_steps"],
              "N_steps_with_data": step_coverage[mid]["n_with_data"],
              "Missing_step_numbers": ",".join(str(x) for x in step_coverage[mid]["missing_steps"])}
             for mid in NOT_RECONSTRUCTABLE]
excl_df = pd.DataFrame(excl_rows)

genome_profile = comp_df.pivot(index="Genome", columns="Module_ID", values="Completeness_%")
genome_profile = genome_profile.loc[genome_ids].reset_index()

unassigned_df = df8[df8["KO_list"].apply(len) == 0][["row", "section", "Gene"]].copy()
outside_catalog_df = unmatched[["row", "section", "Gene", "KO"]].copy() if len(unmatched) else pd.DataFrame(
    columns=["row", "section", "Gene", "KO"])

readme_lines = [
    ("GBM Module Completeness and Statistics (KO-annotated Table 8)", True, 14),
    ("", False, 10),
    (f"Table 8 file: {table8_path}", False, 10),
    (f"Genomes analysed: {N_GENOMES}", False, 10),
    (f"GBM database: GBMs.v1.0 (Valles-Colomer et al., Nat Microbiol 4:623-632, 2019)", False, 10),
    (f"GBM source this run: {GBM_SOURCE_MODE}  |  SHA-256: {GBM_SHA256}", False, 10),
    ("", False, 10),
    ("METHOD", True, 12),
    ("This pipeline uses the KO IDs already assigned in your Table 8 directly — no gene-symbol matching", False, 10),
    ("or module_gene_names.tsv is used or needed. A functional step is present if ANY official OR-", False, 10),
    ("alternative is fully satisfied (ALL KOs of that alternative present; comma-joined KOs = obligate", False, 10),
    ("enzyme-complex subunits, AND-required). A module is scored only if ALL its required steps have at", False, 10),
    ("least one fully data-complete alternative; otherwise it is excluded (Excluded_or_Unresolved sheet),", False, 10),
    ("never estimated. Completeness % = (steps present)/(total required steps) x 100.", False, 10),
    ("Module prevalence: exact (Clopper-Pearson) 95% CI (statsmodels.proportion_confint, method='beta').", False, 10),
    ("Pangenome category: standard Roary/Panaroo convention (Page AJ et al., Bioinformatics 2015):", False, 10),
    ("Core >=99%, Soft-core 95-99%, Shell 15-95%, Cloud <15%.", False, 10),
    ("No per-genome inferential test is applied (see note in Module_Prevalence_CI's source cell).", False, 10),
    ("", False, 10),
    ("SHEETS", True, 12),
    ("- Module_Prevalence_CI: exact 95% CI and pangenome category per reconstructable module.", False, 10),
    ("- Genome_Module_Profile: per-genome completeness (%), wide format.", False, 10),
    ("- Genome_Functional_Steps: per-genome, per-step present/absent (logic-resolved).", False, 10),
    ("- Excluded_or_Unresolved: modules missing step-level KO evidence, and exactly which steps.", False, 10),
    ("- Outside_Official_Catalog: rows whose assigned KO is not part of any official GBM module — check", False, 10),
    ("  these are intentional (e.g. accessory genes tracked for other reasons) rather than a labeling slip.", False, 10),
    ("- Unassigned_Rows: rows with no KO ID at all.", False, 10),
]

with pd.ExcelWriter(CONFIG["output_xlsx"], engine="xlsxwriter") as writer:
    wb = writer.book
    header_fmt = wb.add_format({"bold": True, "bg_color": "#1F4E78", "font_color": "white",
                                 "text_wrap": True, "valign": "vcenter", "border": 1})
    title_fmt = wb.add_format({"bold": True, "font_size": 14})
    sub_fmt = wb.add_format({"bold": True, "font_size": 12})
    normal_fmt = wb.add_format({"font_size": 10})

    ws = wb.add_worksheet("README")
    writer.sheets["README"] = ws
    ws.set_column(0, 0, 110)
    for r, (text, bold, size) in enumerate(readme_lines):
        fmt = title_fmt if (bold and size >= 14) else (sub_fmt if bold else normal_fmt)
        ws.write(r, 0, text, fmt)

    def dump(df, sheet_name):
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        wsx = writer.sheets[sheet_name]
        for c_i, col in enumerate(df.columns):
            wsx.write(0, c_i, col, header_fmt)
            try:
                width = min(60, max(12, df[col].astype(str).str.len().quantile(0.9) + 2))
            except Exception:
                width = 15
            wsx.set_column(c_i, c_i, width)
        wsx.freeze_panes(1, 0)

    dump(prevalence_ci_df, "Module_Prevalence_CI")
    dump(genome_profile, "Genome_Module_Profile")
    dump(steps_df, "Genome_Functional_Steps")
    dump(excl_df, "Excluded_or_Unresolved")
    dump(outside_catalog_df, "Outside_Official_Catalog")
    dump(unassigned_df, "Unassigned_Rows")

print(f"[ok] Workbook written: {CONFIG['output_xlsx']}")



CELL12
# ============================================================
# 12. DOWNLOAD THE RESULT
# ============================================================
if IN_COLAB:
    from google.colab import files
    files.download(CONFIG["output_xlsx"])
else:
    print(f"Output written to: {CONFIG['output_xlsx']} (open it directly).")
